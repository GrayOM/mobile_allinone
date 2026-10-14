from __future__ import annotations

import time
from pathlib import Path

from docx import Document

from backend.app.demo import create_demo_apk


def _wait_for_run(client, run_id: str, timeout: float = 45):
    deadline = time.monotonic() + timeout
    run: dict = {}
    while time.monotonic() < deadline:
        run = client.get(f"/api/runs/{run_id}").json()
        if run["status"] in {"completed", "completed_with_gaps", "failed", "stopped", "manual_required"}:
            return run
        time.sleep(0.05)
    raise AssertionError(f"run did not finish: {run.get('status')} {run.get('current_stage')} {run.get('error')}")


def test_free_assessment_profile_reports_confirmed_findings_only(client, tmp_path: Path):
    project = client.post(
        "/api/projects",
        json={"name": "all-profile", "run_mode": "mock", "assessment_profile": "free_assessment", "ai_enabled": False},
    ).json()
    assert project["assessment_profile"] == "free_assessment"
    apk = create_demo_apk(tmp_path / "all.apk")
    with apk.open("rb") as stream:
        app = client.post(
            f"/api/projects/{project['id']}/apps/upload",
            files={"file": (apk.name, stream, "application/vnd.android.package-archive")},
        ).json()
    plan = client.get(f"/api/apps/{app['id']}/assessment/plan").json()
    assert plan["total"] == 0
    baseline = client.get(
        f"/api/coverage?app_id={app['id']}&standard=free_assessment"
    ).json()
    assert baseline["total_catalog"] == 0
    assert baseline["tests"] == []

    started = client.post(
        "/api/runs",
        json={
            "project_id": project["id"],
            "app_id": app["id"],
            "device_id": "mock-android-01",
            "device_adapter": "mock",
            "proxy_adapter": "mock",
        },
    )
    assert started.status_code == 201
    run = _wait_for_run(client, started.json()["id"])
    assert run["status"] == "completed", run.get("error")
    assert run["options"]["assessment_summary"]["confirmed"] == 0

    # No confirmed vulnerability yet: the DOCX is refused.
    assert client.post(f"/api/runs/{run['id']}/report/docx").status_code == 409

    evidence = client.get(f"/api/runs/{run['id']}/evidence").json()
    storage = next(item for item in evidence if item["evidence_type"] == "storage_snapshot")
    screenshot = next(item for item in evidence if item["evidence_type"] == "screenshot")
    created = client.post(
        "/api/findings",
        json={
            "run_id": run["id"],
            "title": "앱 전용 저장소에 세션 토큰 평문 저장",
            "category": "local_storage",
            "severity": "high",
            "location": "shared_prefs/settings.xml",
            "rationale": "After 스냅샷의 설정 파일에 세션 값이 평문으로 남는다.",
            "reproduction": ["로그인 후 저장소 Before/After 수집", "shared_prefs 변화 확인"],
            "evidence_ids": [screenshot["id"]],
            "reviewer": "claude-code",
        },
    )
    assert created.status_code == 201, created.text
    finding = created.json()
    assert finding["verdict"] == "needs_review"
    assert finding["source"] == "manual"
    assert finding["synthetic"] is True

    # Confirmation requires the category's evidence policy (storage snapshot/diff).
    rejected = client.post(
        f"/api/findings/{finding['id']}/record",
        json={"verdict": "confirmed", "reviewer": "claude-code", "summary": "확정", "criteria_confirmed": True},
    )
    assert rejected.status_code == 422
    assert "storage_snapshot" in rejected.json()["detail"]
    unconfirmed = client.post(
        f"/api/findings/{finding['id']}/record",
        json={
            "verdict": "confirmed",
            "reviewer": "claude-code",
            "summary": "확정",
            "evidence_ids": [storage["id"]],
            "criteria_confirmed": False,
        },
    )
    assert unconfirmed.status_code == 422

    confirmed = client.post(
        f"/api/findings/{finding['id']}/record",
        json={
            "verdict": "confirmed",
            "reviewer": "claude-code",
            "summary": "After 저장소 스냅샷에서 세션 값 평문 저장을 확인했다.",
            "evidence_ids": [storage["id"]],
            "criteria_confirmed": True,
        },
    )
    assert confirmed.status_code == 200, confirmed.text
    assert confirmed.json()["verdict"] == "confirmed"
    sources = client.get(f"/api/findings/{finding['id']}/sources").json()
    source_rows = sources if isinstance(sources, list) else sources.get("sources") or []
    linked = {evidence_id for row in source_rows for evidence_id in (row.get("evidence_ids") or [])}
    assert storage["id"] in linked and screenshot["id"] in linked
    attestations = [
        item
        for item in client.get(f"/api/runs/{run['id']}/evidence").json()
        if item["evidence_type"] == "assessment_attestation"
    ]
    assert len(attestations) == 1 and attestations[0]["finding_id"] == finding["id"]

    refreshed = client.get(f"/api/runs/{run['id']}").json()
    assert refreshed["options"]["assessment_summary"]["confirmed"] == 1
    assert refreshed["options"]["assessment_summary"]["confirmed_manual_findings"] == 1
    assert any(
        item["evidence_type"] == "vulnerability_assessment"
        for item in client.get(f"/api/runs/{run['id']}/evidence").json()
    )

    report = client.post(f"/api/runs/{run['id']}/report/docx")
    assert report.status_code == 200, report.text
    download = client.get(f"/api/runs/{run['id']}/report/docx")
    assert download.status_code == 200
    saved = tmp_path / "report.docx"
    saved.write_bytes(download.content)
    text = "\n".join(paragraph.text for paragraph in Document(saved).paragraphs)
    assert "전체 진단 (ALL)" in text
    assert "앱 전용 저장소에 세션 토큰 평문 저장" in text
    assert "SYNTHETIC MOCK" in text

    # A different manual finding stays out of the report while it is not confirmed.
    other = client.post(
        "/api/findings",
        json={"run_id": run["id"], "title": "검토 중 항목", "category": "navigation", "reviewer": "claude-code"},
    ).json()
    text_after = "\n".join(
        paragraph.text
        for paragraph in Document(
            _save(tmp_path / "report2.docx", client.get(f"/api/runs/{run['id']}/report/docx").content)
        ).paragraphs
    )
    assert "검토 중 항목" not in text_after
    assert other["verdict"] == "needs_review"

    # Downgrading drops the confirmation record and resets confidence.
    downgraded = client.post(
        f"/api/findings/{finding['id']}/record",
        json={"verdict": "false_positive", "reviewer": "claude-code", "summary": "재검토 결과 오탐"},
    ).json()
    assert downgraded["verdict"] == "false_positive"
    assert downgraded["confidence"] == 0.0
    assert downgraded["rationale"] == "After 저장소 스냅샷에서 세션 값 평문 저장을 확인했다."
    sources = client.get(f"/api/findings/{finding['id']}/sources").json()
    linked = {evidence_id for row in sources for evidence_id in (row.get("evidence_ids") or [])}
    assert attestations[0]["id"] not in linked
    assert client.get(f"/api/runs/{run['id']}").json()["options"]["assessment_summary"]["confirmed"] == 0
    assert client.post(f"/api/runs/{run['id']}/report/docx").status_code == 409

    # Ad-hoc categories accept analyst attachments as probative evidence.
    adhoc = client.post(
        "/api/findings",
        json={
            "run_id": run["id"],
            "title": "보안키패드 입력값 화면 노출",
            "category": "secure_keypad",
            "severity": "medium",
            "reviewer": "claude-code",
        },
    ).json()
    png = tmp_path / "keypad.png"
    png.write_bytes(b"\x89PNG\r\n\x1a\nmanual-capture")
    with png.open("rb") as stream:
        attached = client.post(
            f"/api/findings/{adhoc['id']}/evidence",
            files={"file": ("keypad.png", stream, "image/png")},
        )
    assert attached.status_code == 200, attached.text
    assert attached.json()["evidence_type"] == "manual_assessment_attachment"
    with png.open("rb") as stream:
        wrong_kind = client.post(
            f"/api/findings/{adhoc['id']}/evidence",
            files={"file": ("keypad.exe", stream, "application/octet-stream")},
        )
    assert wrong_kind.status_code == 415
    confirmed_adhoc = client.post(
        f"/api/findings/{adhoc['id']}/record",
        json={
            "verdict": "confirmed",
            "reviewer": "claude-code",
            "summary": "키패드 입력 중 캡처한 화면에 입력값이 그대로 보인다.",
            "criteria_confirmed": True,
        },
    )
    assert confirmed_adhoc.status_code == 200, confirmed_adhoc.text
    # The attachment of another finding cannot be reused.
    reused = client.post(
        f"/api/findings/{finding['id']}/record",
        json={
            "verdict": "confirmed",
            "reviewer": "claude-code",
            "summary": "재확정",
            "evidence_ids": [attached.json()["id"], storage["id"]],
            "criteria_confirmed": True,
        },
    )
    assert reused.status_code == 422


def _save(path: Path, data: bytes) -> Path:
    path.write_bytes(data)
    return path


def test_manual_finding_rejects_foreign_evidence_and_running_runs(client, tmp_path: Path):
    demo = client.post("/api/demo/bootstrap").json()
    started = client.post(
        "/api/runs",
        json={
            "project_id": demo["project"]["id"],
            "app_id": demo["app"]["id"],
            "device_id": "mock-android-01",
            "device_adapter": "mock",
            "proxy_adapter": "mock",
        },
    ).json()
    run = _wait_for_run(client, started["id"])
    rejected = client.post(
        "/api/findings",
        json={"run_id": run["id"], "title": "x", "category": "local_storage", "evidence_ids": ["not-an-evidence"]},
    )
    assert rejected.status_code == 422
    missing_run = client.post("/api/findings", json={"run_id": "nope", "title": "x", "category": "local_storage"})
    assert missing_run.status_code == 404
    # Findings created by the domestic ledger are judged through the control endpoint only.
    standard = next(
        (item for item in client.get(f"/api/findings?run_id={run['id']}").json() if item["source"].startswith("standard:")),
        None,
    )
    if standard is not None:
        blocked = client.post(
            f"/api/findings/{standard['id']}/record",
            json={"verdict": "false_positive", "reviewer": "x", "summary": "y"},
        )
        assert blocked.status_code == 409
