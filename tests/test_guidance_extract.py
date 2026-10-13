from __future__ import annotations

from pathlib import Path

import docx

from backend.app.catalog.guidance_extract import build_guidance, parse_advisory, verify_against_catalog


def _advisory(path: Path, titles: list[str]) -> Path:
    document = docx.Document()
    for title in titles:
        document.add_heading(title, level=1)
        document.add_heading("취약점 설명", level=2)
        document.add_paragraph(f"{title} 설명")
        document.add_heading("진단 예시", level=2)
        table = document.add_table(rows=3, cols=2)
        for row, (head, body) in enumerate(
            (("구분", "내용"), ("위협 요소", f"{title} 위협"), ("진단 기준", f"{title} 기준"))
        ):
            table.cell(row, 0).text = head
            table.cell(row, 1).text = body
        document.add_heading("조치 방안", level=2)
        remediation = document.add_table(rows=3, cols=2)
        for row, (head, body) in enumerate(
            (("구분", "내용"), ("조치 방안", f"{title} 조치"), ("Android 소스코드 예시", "code()"))
        ):
            remediation.cell(row, 0).text = head
            remediation.cell(row, 1).text = body
    document.save(str(path))
    return path


def test_extractor_maps_sections_and_numbers_items_in_document_order(tmp_path):
    path = _advisory(tmp_path / "cii.docx", ["코드 인젝션 (Code Injection)", "SQL 인젝션 (SQL Injection)"])
    items = parse_advisory(path, "critical_infrastructure", "CII-MA")
    assert [item["control_id"] for item in items] == ["CII-MA-01", "CII-MA-02"]
    first = items[0]
    assert first["description"] == ["코드 인젝션 (Code Injection) 설명"]
    assert first["threats"] == ["코드 인젝션 (Code Injection) 위협"]
    assert first["diagnostic_criteria"] == ["코드 인젝션 (Code Injection) 기준"]
    assert first["remediation"] == ["코드 인젝션 (Code Injection) 조치"]
    assert first["remediation_examples"] == {"Android 소스코드 예시": ["code()"]}


def test_extractor_refuses_documents_that_do_not_match_the_catalog(tmp_path):
    cii = _advisory(tmp_path / "cii.docx", ["다른 항목"])
    efi = _advisory(tmp_path / "efi.docx", ["[전자금융] 거래 인증수단 검증 오류"])
    problems = verify_against_catalog(build_guidance(cii, efi))
    assert any("CII-MA-01" in line for line in problems)
    assert any("CII-MA-27" in line for line in problems)
