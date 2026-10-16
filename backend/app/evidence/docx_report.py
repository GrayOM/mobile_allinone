from __future__ import annotations

import os
import uuid
from datetime import datetime, timezone
from pathlib import Path
from typing import Any

from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor
from sqlalchemy import select
from sqlalchemy.orm import Session

from backend.app.assessment_scope import custom_row_for, resolve_definition
from backend.app.catalog import execution_plan
from backend.app.catalog.korean_standards import is_free_assessment
from backend.app.catalog.guidance import guidance_for
from backend.app.core.config import AppSettings, get_settings
from backend.app.database.models import (
    FindingSource,
    AppArtifact,
    ControlTest,
    DiagnosticRun,
    Evidence,
    Finding,
    Project,
)


PROFILE_NAMES = {
    "critical_infrastructure": "주요정보통신기반시설",
    "electronic_financial": "전자금융기반시설",
    "free_assessment": "전체 진단 (ALL)",
}
TERMINAL_ASSESSMENT_RESULTS = {"confirmed", "not_vulnerable", "not_applicable"}
EXCLUDED_EVIDENCE_TYPES = {
    "assessment_ledger",
    "assessment_ledger_amendment",
    "assessment_attestation",
    "interaction_checkpoint",
    "vulnerability_assessment",
}


def curate_customer_evidence(items: list[Evidence]) -> list[Evidence]:
    """Keep the analyst's order while removing audit wrappers and duplicate images.

    Raw evidence remains in the Run.  This view is intentionally limited to the
    sources that prove the vulnerability to a report reader; checkpoint and
    attestation containers merely repeat their linked screenshot/log children.
    """
    selected: list[Evidence] = []
    seen_ids: set[str] = set()
    seen_images: set[str] = set()
    for item in items:
        if item.id in seen_ids or item.evidence_type in EXCLUDED_EVIDENCE_TYPES:
            continue
        seen_ids.add(item.id)
        if (item.mime_type or "").startswith("image/"):
            image_key = item.sha256 or str(item.file_path or item.id)
            if image_key in seen_images:
                continue
            seen_images.add(image_key)
        selected.append(item)
    return selected


class VulnerabilityDocxReportRenderer:
    def __init__(self, settings: AppSettings | None = None):
        self.settings = settings or get_settings()

    def render(self, db: Session, run_id: str) -> Path:
        run = db.get(DiagnosticRun, run_id)
        if not run:
            raise LookupError("진단 실행을 찾을 수 없습니다.")
        project = db.get(Project, run.project_id)
        app = db.get(AppArtifact, run.app_id) if run.app_id else None
        if not project or not app:
            raise LookupError("보고서에 필요한 프로젝트 또는 앱을 찾을 수 없습니다.")
        if run.status != "completed":
            raise ValueError(
                "최종 보고서는 진단 Run이 완료된 뒤에만 생성할 수 있습니다. "
                "미점검 항목과 필수 실행 Gap을 먼저 처리하세요."
            )
        profile = str(run.options.get("assessment_profile") or project.assessment_profile)
        if is_free_assessment(profile):
            return self._render_free(db, project=project, app=app, run=run, profile=profile)
        all_controls = db.scalars(
            select(ControlTest)
            .where(
                ControlTest.run_id == run.id,
                ControlTest.standard == profile,
            )
            .order_by(ControlTest.mastg_id)
        ).all()
        unresolved = [
            item
            for item in all_controls
            if item.result not in TERMINAL_ASSESSMENT_RESULTS
        ]
        if unresolved:
            preview = ", ".join(item.mastg_id for item in unresolved[:5])
            suffix = "…" if len(unresolved) > 5 else ""
            raise ValueError(
                f"최종 보고서 생성이 차단됐습니다. 미점검·추가 검토 "
                f"{len(unresolved)}개 항목을 먼저 판정하세요: {preview}{suffix}"
            )
        controls = db.scalars(
            select(ControlTest)
            .where(
                ControlTest.run_id == run.id,
                ControlTest.standard == profile,
                ControlTest.result == "confirmed",
            )
            .order_by(ControlTest.mastg_id)
        ).all()
        if not controls:
            raise ValueError("취약 확정 항목이 없어 DOCX 보고서를 생성하지 않았습니다.")

        finding_ids = {
            finding_id for control in controls for finding_id in control.finding_ids
        }
        findings = {
            item.id: item
            for item in db.scalars(
                select(Finding).where(
                    Finding.id.in_(finding_ids or {"-"}),
                    Finding.run_id == run.id,
                    Finding.verdict == "confirmed",
                )
            ).all()
        }
        evidence_ids = {
            evidence_id for control in controls for evidence_id in control.evidence_ids
        }
        evidence = {
            item.id: item
            for item in db.scalars(
                select(Evidence).where(
                    Evidence.id.in_(evidence_ids or {"-"}),
                    Evidence.run_id == run.id,
                )
            ).all()
        }
        report_evidence = {
            control.id: curate_customer_evidence(
                [evidence[item] for item in control.evidence_ids if item in evidence]
            )
            for control in controls
        }

        document = Document()
        self._configure(document)
        self._cover(document, project, app, run, profile, len(controls))
        self._summary(document, controls, findings, report_evidence)
        document.add_page_break()
        for index, control in enumerate(controls, start=1):
            self._control_section(
                document,
                db=db,
                platform=app.platform,
                index=index,
                control=control,
                findings=[findings[item] for item in control.finding_ids if item in findings],
                evidence=report_evidence[control.id],
            )
            if index != len(controls):
                document.add_page_break()

        return self._save(document, run_id=run.id, profile=profile)

    def _render_free(
        self,
        db: Session,
        *,
        project: Project,
        app: AppArtifact,
        run: DiagnosticRun,
        profile: str,
    ) -> Path:
        """ALL profile: the report is built from confirmed findings of this Run."""
        findings = db.scalars(
            select(Finding)
            .where(Finding.run_id == run.id, Finding.verdict == "confirmed")
            .order_by(Finding.created_at, Finding.id)
        ).all()
        if not findings:
            raise ValueError("취약 확정 발견항목이 없어 DOCX 보고서를 생성하지 않았습니다.")
        sources = db.scalars(
            select(FindingSource).where(
                FindingSource.finding_id.in_([item.id for item in findings])
            )
        ).all()
        linked: dict[str, list[str]] = {item.id: [] for item in findings}
        for source in sources:
            linked[source.finding_id].extend(source.evidence_ids or [])
        for item in findings:
            linked[item.id].extend(evidence.id for evidence in item.evidence)
        evidence_ids = {evidence_id for ids in linked.values() for evidence_id in ids}
        evidence = {
            item.id: item
            for item in db.scalars(
                select(Evidence).where(
                    Evidence.id.in_(evidence_ids or {"-"}),
                    Evidence.run_id == run.id,
                )
            ).all()
        }
        rows_by_finding = {
            finding.id: curate_customer_evidence(
                [
                    evidence[item]
                    for item in dict.fromkeys(linked[finding.id])
                    if item in evidence
                ]
            )
            for finding in findings
        }
        document = Document()
        self._configure(document)
        self._cover(document, project, app, run, profile, len(findings))
        document.add_heading("취약점 요약", level=1)
        table = self._summary_table(document, ("번호", "취약점", "위험도", "증적"))
        for index, finding in enumerate(findings, start=1):
            row = table.add_row().cells
            row[0].text = f"{index:02d}"
            row[1].text = finding.title
            row[2].text = finding.severity.upper()
            row[3].text = f"{len(rows_by_finding[finding.id])}건"
        document.add_page_break()
        for index, finding in enumerate(findings, start=1):
            self._finding_section(
                document,
                platform=app.platform,
                index=index,
                finding=finding,
                evidence=rows_by_finding[finding.id],
            )
            if index != len(findings):
                document.add_page_break()
        return self._save(document, run_id=run.id, profile=profile)

    def _save(self, document: Document, *, run_id: str, profile: str) -> Path:
        output = self.settings.reports_dir / f"{run_id}-{profile}-vulnerabilities.docx"
        temporary = output.with_name(f".{output.name}.{uuid.uuid4()}.tmp.docx")
        try:
            document.save(temporary)
            os.replace(temporary, output)
        finally:
            temporary.unlink(missing_ok=True)
        return output

    @staticmethod
    def _summary_table(document: Document, headers: tuple[str, ...]):
        table = document.add_table(rows=1, cols=len(headers))
        table.style = "Table Grid"
        table.alignment = WD_TABLE_ALIGNMENT.CENTER
        for cell, value in zip(table.rows[0].cells, headers):
            cell.text = value
            cell.paragraphs[0].runs[0].bold = True
            _shade_cell(cell, "17333A")
            cell.paragraphs[0].runs[0].font.color.rgb = RGBColor(255, 255, 255)
        return table

    def _finding_section(
        self,
        document: Document,
        *,
        platform: str,
        index: int,
        finding: Finding,
        evidence: list[Evidence],
    ) -> None:
        document.add_paragraph(f"VULNERABILITY {index:02d}", style="Subtitle")
        document.add_heading(finding.title, level=1)
        facts = document.add_table(rows=2, cols=4)
        facts.style = "Table Grid"
        facts.cell(0, 0).text = "판정"
        facts.cell(0, 1).text = "CONFIRMED"
        facts.cell(0, 2).text = "위험도"
        facts.cell(0, 3).text = finding.severity.upper()
        facts.cell(1, 0).text = "분류"
        facts.cell(1, 1).text = f"{finding.category} · {platform}"
        facts.cell(1, 2).text = "증적 수"
        facts.cell(1, 3).text = str(len(evidence))
        for row in facts.rows:
            for cell_index in (0, 2):
                row.cells[cell_index].paragraphs[0].runs[0].bold = True
                _shade_cell(row.cells[cell_index], "DCE8E6")
        if finding.location:
            document.add_paragraph(f"위치: {finding.location}", style="Caption")
        document.add_heading("판정 근거", level=2)
        document.add_paragraph(finding.rationale or "같은 Run의 재현 증적을 검토자가 확인했습니다.")
        document.add_heading("재현 절차", level=2)
        for step in [item for item in finding.reproduction if item] or ["재현 절차가 기록되지 않았습니다."]:
            document.add_paragraph(step, style="List Number")
        document.add_heading("핵심 재현 증적", level=2)
        if not evidence:
            document.add_paragraph("연결 증적 메타데이터를 찾지 못했습니다.")
        else:
            document.add_paragraph(
                "판정에 사용한 원본 중 고객이 재현 흐름을 확인하는 데 필요한 자료만 "
                "순서대로 표시합니다. 전체 원본과 해시는 작업대 Run에 보존됩니다.",
                style="Caption",
            )
        for position, item in enumerate(evidence, start=1):
            self._evidence(document, position, item)
        if finding.additional_checks:
            document.add_heading("조치 권고", level=2)
            for line in finding.additional_checks:
                document.add_paragraph(line, style="List Bullet")

    @staticmethod
    def _configure(document: Document) -> None:
        section = document.sections[0]
        section.top_margin = Cm(1.8)
        section.bottom_margin = Cm(1.7)
        section.left_margin = Cm(1.8)
        section.right_margin = Cm(1.8)
        normal = document.styles["Normal"]
        normal.font.name = "Malgun Gothic"
        normal._element.rPr.rFonts.set(qn("w:eastAsia"), "맑은 고딕")
        normal.font.size = Pt(9)
        for name, size, color in (
            ("Title", 27, "17333A"),
            ("Heading 1", 18, "17333A"),
            ("Heading 2", 12, "B65722"),
            ("Heading 3", 10, "286A6B"),
        ):
            style = document.styles[name]
            style.font.name = "Malgun Gothic"
            style._element.rPr.rFonts.set(qn("w:eastAsia"), "맑은 고딕")
            style.font.size = Pt(size)
            style.font.color.rgb = RGBColor.from_string(color)

    def _cover(
        self,
        document: Document,
        project: Project,
        app: AppArtifact,
        run: DiagnosticRun,
        profile: str,
        count: int,
    ) -> None:
        marker = document.add_paragraph("MOBILE SECURITY WORKBENCH / CONFIRMED ONLY")
        marker.style = document.styles["Subtitle"]
        marker.runs[0].font.color.rgb = RGBColor(40, 106, 107)
        title = document.add_heading("모바일 앱 취약점 진단 결과보고서", 0)
        title.alignment = WD_ALIGN_PARAGRAPH.LEFT
        subtitle = document.add_paragraph(
            f"{PROFILE_NAMES.get(profile, profile)} · 취약 확정 항목 전용"
        )
        subtitle.runs[0].bold = True
        subtitle.runs[0].font.size = Pt(14)
        subtitle.runs[0].font.color.rgb = RGBColor(182, 87, 34)
        document.add_paragraph("\n")
        table = document.add_table(rows=0, cols=2)
        table.alignment = WD_TABLE_ALIGNMENT.LEFT
        table.style = "Table Grid"
        for label, value in (
            ("프로젝트", project.name),
            ("대상 앱", app.app_name or app.original_name),
            ("Package / Bundle", app.package_name or "미확인"),
            ("앱 SHA-256", app.sha256),
            ("진단 Run", run.id),
            ("취약 확정", f"{count}건"),
            ("생성 시각", datetime.now(timezone.utc).isoformat()),
        ):
            cells = table.add_row().cells
            cells[0].text = label
            cells[1].text = value
            cells[0].vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            cells[0].paragraphs[0].runs[0].bold = True
        document.add_paragraph(
            "\n이 문서는 취약 판정이 확정된 항목과 고객 확인에 필요한 핵심 재현 증적만 "
            "포함합니다. 양호·해당없음·미확정 항목은 수록하지 않습니다. 중복 캡처와 "
            "감사 원장도 제외하며 전체 원본은 Run에 그대로 보존합니다."
        )
        if run.synthetic:
            warning = document.add_paragraph("SYNTHETIC MOCK — 실제 진단 결과로 사용할 수 없습니다.")
            warning.runs[0].bold = True
            warning.runs[0].font.color.rgb = RGBColor(182, 87, 34)

    @staticmethod
    def _summary(
        document: Document,
        controls: list[ControlTest],
        findings: dict[str, Finding],
        report_evidence: dict[str, list[Evidence]],
    ) -> None:
        document.add_heading("취약점 요약", level=1)
        table = VulnerabilityDocxReportRenderer._summary_table(
            document, ("기준 ID", "취약점", "위험도", "증적")
        )
        for control in controls:
            row = table.add_row().cells
            severities = [
                findings[item].severity for item in control.finding_ids if item in findings
            ]
            row[0].text = control.mastg_id
            row[1].text = control.title
            row[2].text = (severities[0] if severities else control.risk).upper()
            row[3].text = f"{len(report_evidence.get(control.id, []))}건"

    def _control_section(
        self,
        document: Document,
        *,
        db: Session,
        platform: str,
        index: int,
        control: ControlTest,
        findings: list[Finding],
        evidence: list[Evidence],
    ) -> None:
        document.add_paragraph(f"VULNERABILITY {index:02d}", style="Subtitle")
        document.add_heading(f"{control.mastg_id}  {control.title}", level=1)
        definition = resolve_definition(
            db,
            project_id=control.project_id,
            standard=control.standard,
            control_id=control.mastg_id,
        )
        plan = execution_plan(definition) if definition else {}
        facts = document.add_table(rows=2, cols=4)
        facts.style = "Table Grid"
        facts.cell(0, 0).text = "판정"
        facts.cell(0, 1).text = "CONFIRMED"
        facts.cell(0, 2).text = "위험도"
        facts.cell(0, 3).text = (findings[0].severity if findings else control.risk).upper()
        facts.cell(1, 0).text = "실행 방식"
        facts.cell(1, 1).text = control.automation.upper()
        facts.cell(1, 2).text = "증적 수"
        facts.cell(1, 3).text = str(len(evidence))
        for row in facts.rows:
            for cell_index in (0, 2):
                row.cells[cell_index].paragraphs[0].runs[0].bold = True
                _shade_cell(row.cells[cell_index], "DCE8E6")

        document.add_heading("판정 근거", level=2)
        document.add_paragraph(control.summary or "같은 Run의 재현 증적과 필수 증적 유형을 확인했습니다.")
        for finding in findings:
            if finding.rationale and finding.rationale != control.summary:
                document.add_paragraph(finding.rationale)

        guidance = guidance_for(control.standard, control.mastg_id)
        custom = custom_row_for(
            db, project_id=control.project_id, control_id=control.mastg_id
        )
        if guidance and guidance["threats"]:
            document.add_heading("위협 요소", level=2)
            for line in guidance["threats"]:
                document.add_paragraph(line, style="List Bullet")

        document.add_heading("취약 판정 기준", level=2)
        for criterion in control.criteria:
            document.add_paragraph(criterion, style="List Bullet")
        if guidance and guidance["diagnostic_criteria"]:
            document.add_paragraph("보안 권고안 진단 기준", style="Caption")
            for line in guidance["diagnostic_criteria"][:20]:
                document.add_paragraph(line, style="List Bullet 2")
        if custom is not None:
            document.add_paragraph(
                f"기준 외 사용자 정의 검사 · 등록 {custom.recorded_by} · 사유: {custom.reason}",
                style="Caption",
            )
            if custom.reference:
                document.add_paragraph(f"참조: {custom.reference}")

        document.add_heading("재현 절차", level=2)
        reproduction = list(
            dict.fromkeys(
                item for finding in findings for item in finding.reproduction if item
            )
        ) or list(custom.procedure if custom is not None and custom.procedure else [])
        reproduction = reproduction or list(control.criteria)
        for step in reproduction:
            document.add_paragraph(step, style="List Number")

        document.add_heading("핵심 재현 증적", level=2)
        if not evidence:
            document.add_paragraph("연결 증적 메타데이터를 찾지 못했습니다.")
        else:
            document.add_paragraph(
                "판정에 사용한 원본 중 고객이 재현 흐름을 확인하는 데 필요한 자료만 "
                "순서대로 표시합니다. 전체 원본과 해시는 작업대 Run에 보존됩니다.",
                style="Caption",
            )
        for position, item in enumerate(evidence, start=1):
            self._evidence(document, position, item)

        document.add_heading("조치 권고", level=2)
        if guidance and guidance["remediation"]:
            for line in guidance["remediation"]:
                document.add_paragraph(line, style="List Bullet")
            platform_key = "Android" if str(platform).endswith("android") else "iOS"
            examples = [
                name
                for name in guidance["remediation_examples"]
                if platform_key.lower() in name.lower()
            ]
            if examples:
                document.add_paragraph(
                    "보안 권고안의 " + ", ".join(examples) + "를 함께 참고합니다.",
                    style="Caption",
                )
        else:
            document.add_paragraph(
                str(plan.get("remediation") or "문서의 보안 권고안에 따라 취약 조건을 제거하고 동일 절차로 재점검합니다.")
            )

    def _evidence(self, document: Document, position: int, item: Evidence) -> None:
        image = self._safe_image_path(item)
        label = "화면" if image else "기술 확인"
        document.add_heading(f"{position}. {label} — {item.title}", level=3)
        if item.description:
            document.add_paragraph(item.description)
        for line in self._evidence_facts(item):
            document.add_paragraph(line, style="List Bullet")
        if image:
            document.add_picture(str(image), width=Cm(15.2))
            document.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER
        integrity = document.add_paragraph(
            f"원본 무결성 · {item.captured_at.isoformat()} · SHA-256 {item.sha256 or 'N/A'}",
            style="Caption",
        )
        integrity.runs[0].font.size = Pt(7)

    @staticmethod
    def _evidence_facts(item: Evidence) -> list[str]:
        """Extract reader-facing facts without dumping commands or raw JSON."""
        data = item.inline_data if isinstance(item.inline_data, dict) else {}
        facts: list[str] = []
        device = data.get("device") if isinstance(data.get("device"), dict) else None
        if device:
            values = [
                str(device.get("model") or "").strip(),
                str(device.get("os_version") or "").strip(),
                "privileged" if device.get("privileged") is True else "",
                str(device.get("frida_status") or "").strip(),
            ]
            summary = " · ".join(value for value in values if value)
            if summary:
                facts.append(f"검증 단말: {summary}")
        for key, label in (
            ("message", "확인 결과"),
            ("script_name", "실행 스크립트"),
            ("category", "검증 범주"),
            ("risk", "승인 위험도"),
            ("content_sha256", "스크립트 SHA-256"),
            ("approved_sha256", "승인 SHA-256"),
            ("status", "수집 상태"),
        ):
            value = data.get(key)
            if isinstance(value, (str, int, float, bool)) and str(value).strip():
                facts.append(f"{label}: {value}")
        operation = data.get("operation")
        if isinstance(operation, dict):
            message = operation.get("message")
            status = operation.get("status")
            if message:
                facts.append(f"확인 결과: {message}")
            elif status:
                facts.append(f"수집 상태: {status}")
        # Keep the report compact. Full command output, syslog and JSON remain
        # available from the Run and are referenced by the integrity line.
        return list(dict.fromkeys(facts))[:6]

    def _safe_image_path(self, item: Evidence) -> Path | None:
        if not item.file_path or not (item.mime_type or "").startswith("image/"):
            return None
        path = Path(item.file_path).resolve()
        root = self.settings.data_dir.resolve()
        if root not in path.parents or not path.is_file() or path.stat().st_size > 10 * 1024 * 1024:
            return None
        return path


def _shade_cell(cell: Any, fill: str) -> None:
    properties = cell._tc.get_or_add_tcPr()
    shading = OxmlElement("w:shd")
    shading.set(qn("w:fill"), fill)
    properties.append(shading)
